"""Manual, local UAE job discovery and application-draft helpers."""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import logging
import os
import re
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "job_hunt_config.json"
PROFILE_PATH = ROOT / "Career-Facts" / "career_profile.json"
TRACKER_PATH = ROOT / "Applications" / "jobs.csv"
EXCEL_PATH = ROOT / "Applications" / "jobs.xlsx"
LOG_PATH = ROOT / "Logs" / "job_hunt.log"
CSV_FIELDS = [
    "job_id", "job_key", "company", "role", "location", "location_priority",
    "salary", "salary_currency", "salary_period", "salary_assessment",
    "salary_monthly_min_aed", "salary_monthly_max_aed", "source_name", "source_url",
    "application_url", "canonical_job_url", "date_found", "last_seen", "match_score",
    "match_reasons", "status", "next_action", "resume_draft", "outreach_draft",
]
UAE_CITIES = {
    "dubai": ("Dubai", 1), "sharjah": ("Sharjah", 2),
    "abu dhabi": ("Abu Dhabi", 3), "ajman": ("Ajman", 4),
    "fujairah": ("Fujairah", 4), "ras al khaimah": ("Ras Al Khaimah", 4),
    "umm al quwain": ("Umm Al Quwain", 4),
}
NON_UAE_COUNTRIES = (
    "india", "pakistan", "saudi arabia", "qatar", "oman", "bahrain",
    "kuwait", "united states", "united kingdom", "singapore",
)
TRACKING_PARAMETERS = {"ref", "source", "gh_src", "lever-source"}
UNSUPPORTED_REQUIREMENTS = {
    "SQL": r"\bsql\b",
    "Oracle Fusion/ERP": r"\boracle\s+(?:fusion|erp)\b",
    "Arabic fluency": r"\b(?:fluent|fluency)\s+in\s+arabic\b",
    "field sales": r"\bfield\s+sales\b",
    "UAE driving license": r"\b(?:valid\s+)?uae\s+driving\s+licen[cs]e\b",
    "software development": r"\bsoftware\s+development\b|\bcoding\s+experience\b",
}


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def load_config() -> dict[str, Any]:
    config = load_json(CONFIG_PATH)
    for source in config.get("sources", []):
        if source.get("provider") not in {"greenhouse", "lever"}:
            raise ValueError(f"Unsupported ATS provider: {source.get('provider')}")
        if not re.fullmatch(r"[A-Za-z0-9_-]+", source.get("board", "")):
            raise ValueError(f"Invalid public ATS board identifier: {source.get('name')}")
    return config


def load_profile() -> dict[str, Any]:
    return load_json(PROFILE_PATH)


def normalize_text(value: Any) -> str:
    text = html.unescape(str(value or "")).casefold()
    text = re.sub(r"<br\s*/?>|</(?:p|li|h[1-6])\s*>", " ", text)
    text = re.sub(r"<[^>]*>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def canonicalize_url(value: str | None) -> str:
    if not value:
        return ""
    parsed = urllib.parse.urlsplit(value.strip())
    if parsed.scheme.casefold() not in {"http", "https"} or not parsed.netloc:
        return ""
    host = parsed.netloc.casefold()
    if host.startswith("www."):
        host = host[4:]
    query = [
        (key, val) for key, val in urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
        if not key.casefold().startswith("utm_") and key.casefold() not in TRACKING_PARAMETERS
    ]
    path = re.sub(r"/{2,}", "/", parsed.path).rstrip("/") or "/"
    return urllib.parse.urlunsplit(
        (parsed.scheme.casefold(), host, path, urllib.parse.urlencode(sorted(query)), "")
    )


def make_job_key(company: str, role: str, location: str, url: str | None) -> str:
    canonical = canonicalize_url(url)
    identity = canonical or "|".join(normalize_text(value) for value in (company, role, location))
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def classify_location(location: str, city_priority: list[str] | None = None) -> tuple[str, int] | None:
    text = normalize_text(location)
    if not text:
        return None
    has_uae = bool(re.search(r"\buae\b|united arab emirates", text))
    if not has_uae and any(country in text for country in NON_UAE_COUNTRIES):
        return None
    priorities = {
        normalize_text(city): index + 1
        for index, city in enumerate(city_priority or ["Dubai", "Sharjah", "Abu Dhabi", "Other UAE"])
    }
    for city, result in UAE_CITIES.items():
        if city in text:
            label = result[0]
            return (label, priorities.get(normalize_text(label), result[1]))
    if has_uae:
        return ("Other UAE", priorities.get("other uae", 4))
    return None


def _amount(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"\d[\d,]*(?:\.\d+)?\s*[kK]?", str(value))
    if not match:
        return None
    token = match.group(0).replace(",", "").replace(" ", "")
    return float(token.rstrip("kK")) * (1000 if token.casefold().endswith("k") else 1)


def _currency(value: str) -> str | None:
    token = value.casefold().replace(".", "").replace(" ", "")
    currencies = {
        "aed": "AED", "دإ": "AED", "usd": "USD", "$": "USD", "eur": "EUR",
        "€": "EUR", "gbp": "GBP", "£": "GBP", "inr": "INR", "₹": "INR",
        "sar": "SAR", "qar": "QAR", "omr": "OMR", "kwd": "KWD", "bhd": "BHD",
    }
    return currencies.get(token)


def _period(value: str) -> str:
    text = normalize_text(value)
    if re.search(r"\b(?:hourly|per hour|/\s*hour|hour|per day|daily|day)\b", text):
        return "non-monthly interval"
    if re.search(r"\b(?:annual|annually|yearly|per year|per annum|p\.a\.|/\s*(?:year|yr)|year|yr)\b", text):
        return "annual"
    if re.search(r"\b(?:monthly|per month|/\s*month|month|mo)\b", text):
        return "monthly"
    return "ambiguous"


def assess_salary(raw_salary: Any, description: str, monthly_target: float) -> dict[str, Any]:
    currency = None
    period = "ambiguous"
    low = high = None
    original = "Undisclosed"
    if isinstance(raw_salary, dict):
        currency = _currency(str(raw_salary.get("currency", "")))
        low, high = _amount(raw_salary.get("min")), _amount(raw_salary.get("max"))
        period = _period(str(raw_salary.get("interval", raw_salary.get("period", ""))))
        if low is not None or high is not None:
            original = str(raw_salary)
    else:
        text = str(raw_salary or "").strip()
        salary_context = True if text else re.search(
            r"salary|compensation|remuneration|pay range", description, re.I
        )
        if not text:
            text = description
        if salary_context:
            cur = r"AED|USD|EUR|GBP|INR|SAR|QAR|OMR|KWD|BHD|د\.?\s*إ\.?|[$€£₹]"
            amount = r"\d[\d,]*(?:\.\d+)?\s*[kK]?"
            ranged = re.search(
                rf"(?:(?P<c1>{cur})\s*)?(?P<low>{amount})\s*(?:-|–|—|to)\s*(?:(?P<c2>{cur})\s*)?(?P<high>{amount})\s*(?P<c3>{cur})?",
                text, re.I,
            )
            single = None if ranged else re.search(
                rf"(?:(?P<c1>{cur})\s*)?(?P<value>{amount})\s*(?P<c2>{cur})?", text, re.I
            )
            match = ranged or single
            if match:
                currency = next(
                    (_currency(match.group(key) or "") for key in ("c1", "c2", "c3")
                     if key in match.groupdict() and match.group(key)), None
                )
                low = _amount(match.group("low")) if ranged else _amount(match.group("value"))
                high = _amount(match.group("high")) if ranged else low
                original, period = match.group(0).strip(), _period(text)

    result = {
        "display": original, "currency": currency or "", "period": period,
        "assessment": "salary undisclosed; verify with employer", "monthly_min_aed": "",
        "monthly_max_aed": "", "exclude": False,
    }
    if low is None and high is None:
        return result
    low = high if low is None else low
    high = low if high is None else high
    if currency != "AED" or period in {"ambiguous", "non-monthly interval"}:
        result["assessment"] = "salary needs human verification; no currency/interval conversion applied"
        return result
    monthly_low, monthly_high = (low, high) if period == "monthly" else (low / 12, high / 12)
    result["monthly_min_aed"], result["monthly_max_aed"] = round(monthly_low, 2), round(monthly_high, 2)
    if monthly_high < monthly_target:
        result["assessment"] = f"below AED {monthly_target:,.0f}/month target"
        result["exclude"] = True
    elif monthly_low >= monthly_target:
        result["assessment"] = f"meets AED {monthly_target:,.0f}/month target"
    else:
        result["assessment"] = f"salary range crosses AED {monthly_target:,.0f}/month target; verify offer floor"
    return result


def role_family(title: str, config: dict[str, Any]) -> dict[str, Any] | None:
    normalized = normalize_text(title)
    matches = []
    for family in config["role_families"]:
        for term in family["title_terms"]:
            if re.search(rf"(?<![a-z0-9]){re.escape(term.casefold())}(?![a-z0-9])", normalized):
                matches.append((len(term), family))
    return max(matches, key=lambda pair: pair[0])[1] if matches else None


def required_years(description: str) -> tuple[int | None, int | None]:
    text = normalize_text(description)
    ranges = [
        (int(m.group(1)), int(m.group(2)))
        for m in re.finditer(r"\b(\d+)\s*(?:-|–|to)\s*(\d+)\s*(?:\+\s*)?(?:years?|yrs?)\b", text)
    ]
    requirements = [int(m.group(1)) for m in re.finditer(r"\b(\d+)\s*\+\s*(?:years?|yrs?)\b", text)]
    requirements.extend(
        int(m.group(1)) for m in re.finditer(r"\b(?:minimum(?: of)?|at least)\s+(\d+)\s+(?:years?|yrs?)\b", text)
    )
    if ranges:
        requirements.extend(low for low, _ in ranges)
    return (max(requirements) if requirements else None, max((high for _, high in ranges), default=None))


def _mandatory_gaps(description: str, profile: dict[str, Any]) -> list[str]:
    text = normalize_text(description)
    documented = normalize_text(" ".join(profile.get("all_skills", [])))
    gaps = []
    for label, pattern in UNSUPPORTED_REQUIREMENTS.items():
        if re.search(pattern, text) and not re.search(pattern, documented):
            for sentence in re.split(r"[.;\n]", text):
                if re.search(pattern, sentence) and re.search(r"\b(?:required|must|essential|minimum|need)\b", sentence):
                    gaps.append(label)
                    break
    return gaps


def score_job(job: dict[str, Any], config: dict[str, Any], profile: dict[str, Any]) -> tuple[int, list[str], str]:
    title, description = job.get("title", ""), normalize_text(job.get("description", ""))
    normalized_title = normalize_text(title)
    for term in config["seniority_exclusions"]:
        if re.search(rf"(?<![a-z0-9]){re.escape(term.casefold())}(?![a-z0-9])", normalized_title):
            return 0, [], f"excluded seniority: {term}"
    family = role_family(title, config)
    if not family:
        return 0, [], "title is outside configured target role families"
    score, reasons = 30, [f"target role family: {family['name']}"]
    profile_groups = profile.get("evidence_groups", {})
    matched_groups, matched_terms = [], []
    for group in family["evidence_groups"]:
        if group not in profile_groups:
            continue
        hits = [term for term in config["evidence_terms"].get(group, []) if term.casefold() in description]
        if hits:
            matched_groups.append(group)
            matched_terms.extend(hits[:2])
            score += 12
    if matched_groups:
        reasons.append("documented evidence overlaps: " + ", ".join(dict.fromkeys(matched_terms[:6])))
    else:
        reasons.append("no direct overlap found between requirements and documented evidence")

    min_years, max_years = required_years(description)
    candidate_years = int(profile.get("experience_years", 0))
    if min_years is not None:
        if min_years > 5:
            return 0, reasons, f"excluded experience requirement: {min_years}+ years exceeds documented experience"
        if min_years <= candidate_years:
            score += 10 if max_years is None or max_years >= candidate_years else 6
            reasons.append(f"experience requirement is consistent with documented {candidate_years} years")
        else:
            score -= 10
            reasons.append(f"ad requests {min_years}+ years; resume documents {candidate_years}")
    if "bachelor" in description and profile.get("education"):
        score += 4
        reasons.append("bachelor-level education is documented")
    if any(word in description for word in ("excel", "google sheets")) and any(
        word in normalize_text(" ".join(profile.get("all_skills", []))) for word in ("excel", "google sheets")
    ):
        score += 4
        reasons.append("spreadsheet tools in the ad match documented skills")
    gaps = _mandatory_gaps(description, profile)
    if gaps:
        score -= min(24, 12 * len(gaps))
        reasons.append("required skills not evidenced in resume: " + ", ".join(gaps))
    return max(0, min(100, score)), reasons, ""


def listing_is_expired(raw: dict[str, Any], today: date | None = None) -> bool:
    today = today or date.today()
    if str(raw.get("status", "")).casefold() in {"closed", "expired", "inactive", "filled", "archived"}:
        return True
    deadline = raw.get("application_deadline") or raw.get("validThrough") or raw.get("expiresAt")
    if not deadline:
        return False
    try:
        deadline_date = datetime.fromisoformat(str(deadline).replace("Z", "+00:00")).date()
    except ValueError:
        try:
            deadline_date = date.fromisoformat(str(deadline)[:10])
        except ValueError:
            return False
    return deadline_date < today


def source_endpoint(source: dict[str, Any]) -> str:
    board = urllib.parse.quote(source["board"], safe="")
    if source["provider"] == "greenhouse":
        return f"https://boards-api.greenhouse.io/v1/boards/{board}/jobs?content=true"
    return f"https://api.lever.co/v0/postings/{board}?mode=json"


def fetch_source(source: dict[str, Any], timeout: int = 20) -> list[dict[str, Any]]:
    request = urllib.request.Request(
        source_endpoint(source),
        headers={"User-Agent": "JobHuntSystem/1.0 (manual public job discovery)", "Accept": "application/json"},
        method="GET",
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = json.loads(response.read().decode("utf-8"))
    jobs = payload.get("jobs", []) if source["provider"] == "greenhouse" else payload
    if not isinstance(jobs, list):
        raise ValueError(f"Unexpected response shape from {source['name']}")
    return jobs


def normalize_listing(raw: dict[str, Any], source: dict[str, Any]) -> dict[str, Any]:
    if source["provider"] == "greenhouse":
        location_value = raw.get("location", "")
        location = location_value.get("name", "") if isinstance(location_value, dict) else str(location_value)
        url = raw.get("absolute_url", "")
        description = raw.get("content", "") or ""
        salary_raw = raw.get("salary") or raw.get("compensation")
        company, role = raw.get("company_name") or source["name"], raw.get("title", "")
    else:
        categories = raw.get("categories") or {}
        location, url = categories.get("location", ""), raw.get("hostedUrl", "")
        description = raw.get("descriptionPlain") or raw.get("description", "") or ""
        salary_raw = raw.get("salaryRange") or raw.get("salary")
        company, role = raw.get("company") or source["name"], raw.get("text", "")
    return {
        "company": str(company).strip(), "title": str(role).strip(), "location": str(location).strip(),
        "application_url": raw.get("applyUrl", "") or url, "canonical_job_url": canonicalize_url(url),
        "source_name": source["name"], "source_url": source_endpoint(source),
        "description": html.unescape(str(description)), "salary_raw": salary_raw, "raw": raw,
    }


def evaluate_listing(listing: dict[str, Any], config: dict[str, Any], profile: dict[str, Any], today: date | None = None):
    raw = listing["raw"]
    if listing_is_expired(raw, today):
        return None, "expired"
    location = classify_location(listing["location"], config.get("city_priority"))
    if not location:
        return None, "outside UAE or location unverified"
    salary = assess_salary(listing["salary_raw"], listing["description"], float(config["salary_target"]["monthly_min"]))
    if salary["exclude"]:
        return None, "below salary target"
    score, reasons, exclusion = score_job(listing, config, profile)
    if exclusion:
        return None, exclusion
    if score < int(config["minimum_relevance_score"]):
        return None, "relevance score below minimum"
    job_key = make_job_key(listing["company"], listing["title"], listing["location"], listing["canonical_job_url"])
    today_string = (today or date.today()).isoformat()
    reasons += [f"location priority: {location[0]}", salary["assessment"]]
    if score >= int(config["review_threshold"]):
        next_action = "Review listing, salary, and application requirements"
    elif "verify" in salary["assessment"] or "crosses" in salary["assessment"]:
        next_action = "Review role fit and verify monthly salary with employer"
    else:
        next_action = "Review match and listing details"
    return ({
        "job_id": job_key[:12], "job_key": job_key, "company": listing["company"], "role": listing["title"],
        "location": listing["location"], "location_priority": location[1], "salary": salary["display"],
        "salary_currency": salary["currency"], "salary_period": salary["period"],
        "salary_assessment": salary["assessment"], "salary_monthly_min_aed": salary["monthly_min_aed"],
        "salary_monthly_max_aed": salary["monthly_max_aed"], "source_name": listing["source_name"],
        "source_url": listing["source_url"], "application_url": listing["application_url"],
        "canonical_job_url": listing["canonical_job_url"], "date_found": today_string,
        "last_seen": today_string, "match_score": score, "match_reasons": "; ".join(reasons),
        "status": "New", "next_action": next_action,
    }, "included")


def read_tracker(path: Path = TRACKER_PATH) -> list[dict[str, str]]:
    if not path.exists() or path.stat().st_size == 0:
        return []
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def upsert_tracker(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [dict(row) for row in existing]
    index = {}
    for position, row in enumerate(rows):
        key = row.get("job_key") or make_job_key(
            row.get("company", ""), row.get("role", ""), row.get("location", ""),
            row.get("canonical_job_url") or row.get("application_url"),
        )
        index.setdefault(key, position)
    for new_row in incoming:
        key = new_row["job_key"]
        if key in index:
            current = rows[index[key]]
            for field, value in new_row.items():
                if field not in {"status", "next_action", "date_found"}:
                    current[field] = value
            current.setdefault("status", "New")
            current.setdefault("next_action", "Review listing")
        else:
            rows.append(dict(new_row))
            index[key] = len(rows) - 1
    return rows


def write_tracker(rows: list[dict[str, Any]], path: Path = TRACKER_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(
            "w", newline="", encoding="utf-8", dir=path.parent,
            prefix=f"{path.name}.", suffix=".tmp", delete=False,
        ) as handle:
            temp_name = handle.name
            writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows({field: row.get(field, "") for field in CSV_FIELDS} for row in rows)
        os.replace(temp_name, path)
    finally:
        if temp_name and os.path.exists(temp_name):
            os.unlink(temp_name)


def export_tracker_xlsx(rows: list[dict[str, Any]], path: Path = EXCEL_PATH) -> None:
    """Write a readable Excel tracker with filterable headers and clickable job links."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "UAE Job Matches"
    sheet.append(CSV_FIELDS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sheet.append([row.get(field, "") for field in CSV_FIELDS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row_index in range(2, sheet.max_row + 1):
        for column_name in ("application_url", "canonical_job_url", "source_url"):
            column_index = CSV_FIELDS.index(column_name) + 1
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(cell.value, str) and cell.value.startswith(("https://", "http://")):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    for column_index, field in enumerate(CSV_FIELDS, start=1):
        values = [str(sheet.cell(row=r, column=column_index).value or "") for r in range(1, min(sheet.max_row, 150) + 1)]
        max_length = max((len(value) for value in values), default=len(field))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 48)
    workbook.save(path)


def setup_logging() -> None:
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(filename=LOG_PATH, level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8")


def discover(dry_run: bool = False, skip_drafts: bool = False):
    config, profile = load_config(), load_profile()
    incoming, counts, errors = [], {}, []
    for source in config["sources"]:
        if not source.get("enabled", False):
            continue
        try:
            raw_jobs = fetch_source(source, int(config.get("request_timeout_seconds", 20)))
            counts[source["name"]] = len(raw_jobs)
            logging.info("Fetched %d listings from %s", len(raw_jobs), source["name"])
            for raw in raw_jobs:
                listing = normalize_listing(raw, source)
                row, reason = evaluate_listing(listing, config, profile)
                if row:
                    incoming.append(row)
                else:
                    logging.debug("Skipped %s: %s", raw.get("title", "untitled"), reason)
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, ValueError) as error:
            message = f"{source['name']}: {type(error).__name__}: {error}"
            errors.append(message)
            logging.error("Source failed: %s", message)
    incoming.sort(key=lambda row: (int(row["location_priority"]), -int(row["match_score"]), row["company"].casefold()))
    if not dry_run:
        existing = read_tracker()
        if incoming and not skip_drafts:
            incoming = generate_review_drafts(incoming, existing, config, profile)
        merged = upsert_tracker(existing, incoming)
        write_tracker(merged)
        export_tracker_xlsx(merged)
        logging.info("Tracker upsert complete; eligible incoming listings: %d", len(incoming))
    return incoming, counts, errors


def extract_docx_paragraphs(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
    root = ET.fromstring(document_xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = [
        "".join(node.text or "" for node in paragraph.iter(namespace + "t")).strip()
        for paragraph in root.iter(namespace + "p")
    ]
    return [paragraph for paragraph in paragraphs if paragraph]


def _safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")[:64] or "job"


def _ranked_highlights(job: dict[str, Any], profile: dict[str, Any], config: dict[str, Any]) -> list[str]:
    family = role_family(job.get("role", ""), config)
    groups = set(family["evidence_groups"] if family else [])
    ranked = []
    position = 0
    for experience in profile["experience"]:
        for item in experience["highlights"]:
            overlap = len(groups.intersection(item["groups"]))
            if overlap:
                ranked.append((-overlap, position, item["text"]))
            position += 1
    ranked.sort()
    return [item[2] for item in ranked]


def build_drafts(job: dict[str, Any], profile: dict[str, Any], config: dict[str, Any]) -> tuple[str, str]:
    if not extract_docx_paragraphs(ROOT / profile["source_master_resume"]):
        raise ValueError("Approved master resume contains no readable text")
    family = role_family(job.get("role", ""), config)
    groups = set(family["evidence_groups"] if family else [])
    skills = [
        f"{category}: {', '.join(values)}" for category, values in profile["skills"].items()
        if set(profile["skill_groups"][category]).intersection(groups)
    ]
    experience_lines = ["## Relevant professional experience"]
    highlights = []
    for experience in profile["experience"]:
        relevant = [item["text"] for item in experience["highlights"] if groups.intersection(item["groups"])]
        if relevant:
            experience_lines.extend([
                f"### {experience['title']}",
                f"{experience['company']} | {experience['location']} | {experience['dates']}",
                *[f"- {item}" for item in relevant],
                "",
            ])
            highlights.extend(relevant)
    highlights = list(dict.fromkeys(highlights))[:6]
    resume = "\n".join([
        f"# {profile['name']}", profile["headline"], profile["location"], "",
        "## Targeted profile", profile["summary"], "", "## Relevant skills",
        *[f"- {skill}" for skill in skills], "", *experience_lines, "## Education and credentials",
        *[f"- {item}" for item in profile["education"]],
        *[f"- {item}" for item in profile["certifications"]], "", "## Languages",
        ", ".join(profile["languages"]), "",
        f"Draft for review: {job['role']} at {job['company']}.",
        f"Facts transcribed from approved master resume: {profile['source_master_resume']}.",
    ])
    outreach_evidence = "\n".join(f"- {item}" for item in _ranked_highlights(job, profile, config)[:2])
    outreach = (
        "Draft for review only - not sent.\n\n"
        f"Subject: {job['role']} application\n\n"
        "Hello Hiring Team,\n\n"
        f"I am interested in the {job['role']} opportunity at {job['company']}. "
        "My background includes UAE-based B2B SaaS implementation, client onboarding, and product/operations work. "
        f"Relevant experience includes:\n{outreach_evidence}\n\n"
        "I would welcome a conversation about how this experience may fit your team.\n\n"
        f"Regards,\n{profile['name']}\n"
    )
    return resume, outreach


def write_unique_draft(job: dict[str, Any], profile: dict[str, Any], config: dict[str, Any]) -> tuple[Path, Path]:
    resume_text, outreach_text = build_drafts(job, profile, config)
    draft_dir = ROOT / "Outreach" / "Drafts"
    draft_dir.mkdir(parents=True, exist_ok=True)
    base, suffix, number = _safe_filename(f"{job['company']}_{job['role']}_{job['job_id']}"), date.today().strftime("%Y%m%d"), 1
    while True:
        extra = "" if number == 1 else f"_{number:02d}"
        resume_path = draft_dir / f"{suffix}_{base}{extra}_resume.md"
        outreach_path = draft_dir / f"{suffix}_{base}{extra}_outreach.md"
        if not resume_path.exists() and not outreach_path.exists():
            break
        number += 1
    resume_path.write_text(resume_text, encoding="utf-8")
    outreach_path.write_text(outreach_text, encoding="utf-8")
    return resume_path, outreach_path


def should_generate_review_draft(job: dict[str, Any], existing: dict[str, Any] | None, threshold: int) -> bool:
    already_generated = existing and existing.get("resume_draft") and existing.get("outreach_draft")
    return int(job.get("match_score") or 0) >= threshold and not already_generated


def generate_review_drafts(
    incoming: list[dict[str, Any]],
    existing_rows: list[dict[str, Any]],
    config: dict[str, Any],
    profile: dict[str, Any],
) -> list[dict[str, Any]]:
    existing_by_key = {row.get("job_key"): row for row in existing_rows}
    for job in incoming:
        previous = existing_by_key.get(job["job_key"])
        if should_generate_review_draft(job, previous, int(config["review_threshold"])):
            resume_path, outreach_path = write_unique_draft(job, profile, config)
            job["resume_draft"] = str(resume_path.relative_to(ROOT))
            job["outreach_draft"] = str(outreach_path.relative_to(ROOT))
        elif previous:
            job["resume_draft"] = previous.get("resume_draft", "")
            job["outreach_draft"] = previous.get("outreach_draft", "")
        existing_by_key[job["job_key"]] = job
    return incoming


def run_cli(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Local UAE job discovery; no applications or messages are submitted.")
    commands = parser.add_subparsers(dest="command", required=True)
    discover_parser = commands.add_parser("discover", help="Fetch configured public ATS feeds and update the CSV tracker")
    discover_parser.add_argument("--dry-run", action="store_true", help="Fetch and score listings without changing the tracker")
    discover_parser.add_argument("--skip-drafts", action="store_true", help="Update trackers without creating resume/outreach drafts")
    list_parser = commands.add_parser("list", help="Show tracked jobs")
    list_parser.add_argument("--all", action="store_true", help="Include jobs below the review threshold")
    draft_parser = commands.add_parser("draft", help="Create resume and outreach drafts for a tracked job")
    draft_parser.add_argument("job_id", help="The job_id from Applications/jobs.csv")
    args = parser.parse_args(argv)
    setup_logging()
    try:
        if args.command == "discover":
            incoming, counts, errors = discover(args.dry_run, args.skip_drafts)
            for source, count in counts.items():
                print(f"{source}: fetched {count} public listings")
            for row in incoming:
                print(f"{row['location']:<24} {row['match_score']:>3}  {row['company']} | {row['role']} | {row['salary_assessment']}")
            suffix = " (dry run; tracker unchanged)" if args.dry_run else ""
            print(f"Eligible UAE listings: {len(incoming)}{suffix}")
            if not args.dry_run:
                print(f"Excel tracker: {EXCEL_PATH.relative_to(ROOT)}")
            for error in errors:
                print(f"Source error: {error}")
            return 1 if errors and not counts else 0
        rows = read_tracker()
        if args.command == "list":
            config = load_config()
            for row in rows:
                if not args.all and int(row.get("match_score") or 0) < int(config["review_threshold"]):
                    continue
                print(f"{row.get('job_id')} | {row.get('match_score')} | {row.get('company')} | {row.get('role')} | {row.get('location')} | {row.get('salary_assessment')} | {row.get('status')}")
            return 0
        job = next((row for row in rows if row.get("job_id") == args.job_id or row.get("job_key") == args.job_id), None)
        if not job:
            print(f"No tracked job matches {args.job_id}")
            return 2
        resume_path, outreach_path = write_unique_draft(job, load_profile(), load_config())
        job["resume_draft"] = str(resume_path.relative_to(ROOT))
        job["outreach_draft"] = str(outreach_path.relative_to(ROOT))
        write_tracker(rows)
        print(f"Resume draft: {resume_path.relative_to(ROOT)}")
        print(f"Outreach draft: {outreach_path.relative_to(ROOT)}")
        print("Drafts were saved locally only; nothing was sent or submitted.")
        return 0
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, ET.ParseError) as error:
        logging.exception("Command failed")
        print(f"Error: {error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(run_cli())