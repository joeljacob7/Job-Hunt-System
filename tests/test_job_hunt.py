import unittest
import tempfile
import json
from io import BytesIO
from pathlib import Path
from datetime import date
from unittest.mock import patch

import job_hunt


class JobHuntTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = job_hunt.load_config()
        cls.profile = job_hunt.load_profile()

    def test_excel_export_has_clickable_application_link(self):
        from openpyxl import load_workbook
        row = {field: "" for field in job_hunt.CSV_FIELDS}
        row.update({"job_id": "abc123", "company": "Example", "role": "Implementation Specialist", "application_url": "https://example.com/apply"})
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "jobs.xlsx"
            job_hunt.export_tracker_xlsx([row], output)
            workbook = load_workbook(output)
            sheet = workbook["UAE Job Matches"]
            self.assertEqual(sheet.cell(2, job_hunt.CSV_FIELDS.index("company") + 1).value, "Example")
            link_cell = sheet.cell(2, job_hunt.CSV_FIELDS.index("application_url") + 1)
            self.assertEqual(link_cell.hyperlink.target, "https://example.com/apply")

    def test_canonical_urls_deduplicate_tracking_variants(self):
        first = "https://www.boards.greenhouse.io/careem/jobs/123?utm_source=mail&gh_jid=123#apply"
        second = "https://boards.greenhouse.io/careem/jobs/123?gh_jid=123"
        self.assertEqual(job_hunt.canonicalize_url(first), job_hunt.canonicalize_url(second))
        self.assertEqual(job_hunt.make_job_key("Careem", "Account Manager", "Dubai", first), job_hunt.make_job_key("Careem", "Account Manager", "Dubai", second))

    def test_provider_endpoints_include_new_public_boards_without_changing_existing_urls(self):
        sources = {source["name"]: source for source in self.config["sources"]}
        self.assertEqual(
            job_hunt.source_endpoint(sources["Careem"]),
            "https://boards-api.greenhouse.io/v1/boards/careem/jobs?content=true",
        )
        self.assertEqual(
            job_hunt.source_endpoint(sources["Deliverect"]),
            "https://api.lever.co/v0/postings/deliverect?mode=json",
        )
        self.assertEqual(
            job_hunt.source_endpoint(sources["Ziina"]),
            "https://api.ashbyhq.com/posting-api/job-board/ziina?includeCompensation=true",
        )
        self.assertEqual(
            job_hunt.source_endpoint(sources["Tabby"]),
            "https://tabby.pinpointhq.com/postings.json",
        )
        self.assertEqual(
            job_hunt.source_endpoint({"provider": "workable", "board": "bayutdubizzle"}),
            "https://apply.workable.com/api/v1/widget/accounts/bayutdubizzle?details=true",
        )

    def test_new_provider_feeds_extract_their_documented_job_arrays(self):
        fixtures = (
            (
                {"provider": "ashby", "board": "ziina", "name": "Ziina"},
                {"jobs": [{"title": "Role"}, {"title": "Unlisted", "isListed": False}]},
            ),
            ({"provider": "pinpoint", "board": "tabby", "name": "Tabby"}, {"data": [{"title": "Role"}]}),
            (
                {"provider": "workable", "board": "bayutdubizzle", "name": "Bayut | dubizzle"},
                {"name": "Bayut | dubizzle", "jobs": [{"title": "Role"}]},
            ),
        )
        for source, payload in fixtures:
            with self.subTest(provider=source["provider"]):
                response = BytesIO(json.dumps(payload).encode("utf-8"))
                with patch("job_hunt.urllib.request.urlopen", return_value=response):
                    self.assertEqual(job_hunt.fetch_source(source), [{"title": "Role"}])

    def test_new_provider_feeds_reject_malformed_responses(self):
        fixtures = (
            ({"provider": "ashby", "board": "ziina", "name": "Ziina"}, {"jobs": {}}),
            ({"provider": "pinpoint", "board": "tabby", "name": "Tabby"}, {"data": [{"title": "Role"}, None]}),
            (
                {"provider": "workable", "board": "bayutdubizzle", "name": "Bayut | dubizzle"},
                {"name": "Bayut | dubizzle", "jobs": [{"title": "Role"}, None]},
            ),
        )
        for source, payload in fixtures:
            with self.subTest(provider=source["provider"]):
                response = BytesIO(json.dumps(payload).encode("utf-8"))
                with patch("job_hunt.urllib.request.urlopen", return_value=response):
                    with self.assertRaisesRegex(ValueError, "Unexpected response shape"):
                        job_hunt.fetch_source(source)

    def test_workable_normalization_maps_fields_and_undisclosed_salary(self):
        source = {"provider": "workable", "board": "bayutdubizzle", "name": "Bayut | dubizzle"}
        listing = job_hunt.normalize_listing({
            "title": " Account Manager - Partner Success & Sales ",
            "locations": [{
                "city": "Dubai", "region": "Dubai",
                "country": "United Arab Emirates", "hidden": False,
            }],
            "url": "https://apply.workable.com/j/3165D44D54?utm_source=careers",
            "application_url": "https://apply.workable.com/j/3165D44D54/apply",
            "description": "<p>Manage partner relationships and onboard customers.</p>",
        }, source)
        self.assertEqual(listing["title"], "Account Manager - Partner Success & Sales")
        self.assertEqual(listing["location"], "Dubai, Dubai, United Arab Emirates")
        self.assertEqual(listing["description"], "<p>Manage partner relationships and onboard customers.</p>")
        self.assertEqual(listing["application_url"], "https://apply.workable.com/j/3165D44D54/apply")
        self.assertEqual(listing["canonical_job_url"], "https://apply.workable.com/j/3165D44D54")
        salary = job_hunt.assess_salary(
            listing["salary_raw"], listing["description"], self.config["salary_target"]["monthly_min"]
        )
        self.assertFalse(salary["exclude"])
        self.assertIn("undisclosed", salary["assessment"])
        self.assertIn("verify", salary["assessment"])

    def test_workable_eligibility_keeps_uae_filter_and_unknown_salary_verification(self):
        source = {"provider": "workable", "board": "bayutdubizzle", "name": "Bayut | dubizzle"}
        raw = {
            "title": "Account Manager - Partner Success & Sales",
            "locations": [{
                "city": "Dubai", "region": "Dubai",
                "country": "United Arab Emirates", "hidden": False,
            }],
            "description": "3+ years of account management, partner management, customer success, client training, and customer onboarding.",
            "url": "https://apply.workable.com/j/3165D44D54",
            "application_url": "https://apply.workable.com/j/3165D44D54/apply",
        }
        row, reason = job_hunt.evaluate_listing(
            job_hunt.normalize_listing(raw, source), self.config, self.profile
        )
        self.assertIsNotNone(row, reason)
        self.assertIn("verify", row["salary_assessment"])
        self.assertEqual(row["salary_monthly_min_aed"], "")

        outside = job_hunt.normalize_listing({
            **raw,
            "locations": [{
                "city": "Riyadh", "region": "Riyadh",
                "country": "Saudi Arabia", "hidden": False,
            }],
        }, source)
        row, reason = job_hunt.evaluate_listing(outside, self.config, self.profile)
        self.assertIsNone(row)
        self.assertEqual(reason, "outside UAE or location unverified")

    def test_workable_canonical_urls_deduplicate_tracking_variants(self):
        source = {"provider": "workable", "board": "bayutdubizzle", "name": "Bayut | dubizzle"}
        base = {
            "title": "Account Manager",
            "city": "Dubai",
            "country": "United Arab Emirates",
            "description": "Manage customer accounts.",
        }
        first = job_hunt.normalize_listing({
            **base, "url": "https://apply.workable.com/j/3165D44D54?utm_source=email",
        }, source)
        second = job_hunt.normalize_listing({
            **base, "url": "https://apply.workable.com/j/3165D44D54",
        }, source)
        self.assertEqual(
            job_hunt.make_job_key(first["company"], first["title"], first["location"], first["canonical_job_url"]),
            job_hunt.make_job_key(second["company"], second["title"], second["location"], second["canonical_job_url"]),
        )

    def test_ashby_normalization_and_salary_use_existing_monthly_target_logic(self):
        source = {"provider": "ashby", "board": "ziina", "name": "Ziina"}
        listing = job_hunt.normalize_listing({
            "title": "Product Operations Associate",
            "location": "Dubai",
            "descriptionPlain": "3+ years in product operations. Own UAT, workflows, dashboards, and Excel reporting.",
            "jobUrl": "https://jobs.ashbyhq.com/ziina/role-1",
            "applyUrl": "https://jobs.ashbyhq.com/ziina/role-1/application",
            "compensation": {
                "scrapeableCompensationSalarySummary": "AED 78,000-96,000 per year",
            },
        }, source)
        self.assertEqual(listing["company"], "Ziina")
        self.assertEqual(listing["location"], "Dubai")
        self.assertEqual(listing["application_url"], "https://jobs.ashbyhq.com/ziina/role-1/application")
        salary = job_hunt.assess_salary(
            listing["salary_raw"], listing["description"], self.config["salary_target"]["monthly_min"]
        )
        self.assertEqual(salary["monthly_min_aed"], 6500)
        self.assertEqual(salary["assessment"], "meets AED 6,500/month target")

    def test_pinpoint_normalization_maps_location_salary_description_and_deadline(self):
        source = {"provider": "pinpoint", "board": "tabby", "name": "Tabby"}
        listing = job_hunt.normalize_listing({
            "title": "Customer Success Associate",
            "url": "https://tabby.pinpointhq.com/en/postings/role-1",
            "description": "Support customer onboarding and adoption.",
            "benefits": "",
            "location": {"city": "Dubai", "province": "Dubai", "name": "United Arab Emirates"},
            "compensation_minimum": 6500,
            "compensation_maximum": 8000,
            "compensation_currency": "AED",
            "compensation_frequency": "monthly",
            "deadline_at": "2026-12-31T23:59:59+04:00",
        }, source)
        self.assertIn("Dubai", listing["location"])
        self.assertIn("United Arab Emirates", listing["location"])
        self.assertEqual(listing["source_url"], "https://tabby.pinpointhq.com/postings.json")
        self.assertEqual(listing["raw"]["application_deadline"], "2026-12-31T23:59:59+04:00")
        salary = job_hunt.assess_salary(
            listing["salary_raw"], listing["description"], self.config["salary_target"]["monthly_min"]
        )
        self.assertEqual(salary["assessment"], "meets AED 6,500/month target")

    def test_pinpoint_undisclosed_salary_remains_verification_required(self):
        listing = job_hunt.normalize_listing({
            "title": "Business Analyst",
            "description": "Competitive compensation for the right candidate.",
            "location": {"city": "Dubai", "name": "UAE"},
        }, {"provider": "pinpoint", "board": "tabby", "name": "Tabby"})
        salary = job_hunt.assess_salary(
            listing["salary_raw"], listing["description"], self.config["salary_target"]["monthly_min"]
        )
        self.assertFalse(salary["exclude"])
        self.assertIn("undisclosed", salary["assessment"])
        self.assertIn("verify", salary["assessment"])

    def test_pinpoint_foreign_salary_from_benefits_is_flagged_without_conversion(self):
        listing = job_hunt.normalize_listing({
            "title": "Business Analyst",
            "description": "Analyze business operations.",
            "benefits": "<p>Base salary: USD 90,000 per year.</p>",
            "location": {"city": "Dubai", "name": "UAE"},
        }, {"provider": "pinpoint", "board": "tabby", "name": "Tabby"})
        salary = job_hunt.assess_salary(
            listing["salary_raw"], listing["description"], self.config["salary_target"]["monthly_min"]
        )
        self.assertFalse(salary["exclude"])
        self.assertEqual(salary["currency"], "USD")
        self.assertIn("verification", salary["assessment"])

    def test_new_provider_location_filtering_and_relevance_are_preserved(self):
        source = {"provider": "pinpoint", "board": "tabby", "name": "Tabby"}
        raw = {
            "title": "Product Operations Associate",
            "url": "https://tabby.pinpointhq.com/en/postings/product-ops",
            "description": "3+ years in product operations. Own UAT, feature definition, workflows, dashboards, and Excel reporting.",
            "location": {"city": "Dubai", "province": "Dubai", "name": "UAE"},
        }
        listing = job_hunt.normalize_listing(raw, source)
        row, reason = job_hunt.evaluate_listing(listing, self.config, self.profile)
        self.assertIsNotNone(row, reason)
        self.assertGreaterEqual(row["match_score"], self.config["minimum_relevance_score"])
        outside = job_hunt.normalize_listing({
            **raw,
            "location": {"city": "Riyadh", "province": "Riyadh", "name": "Saudi Arabia"},
        }, source)
        self.assertIsNone(job_hunt.evaluate_listing(outside, self.config, self.profile)[0])

    def test_new_provider_job_urls_deduplicate_tracking_variants(self):
        source = {"provider": "ashby", "board": "ziina", "name": "Ziina"}
        base = {
            "title": "Product Operations Associate",
            "location": "Dubai",
            "descriptionPlain": "Product operations",
        }
        first = job_hunt.normalize_listing({
            **base, "jobUrl": "https://jobs.ashbyhq.com/ziina/123?utm_source=email",
        }, source)
        second = job_hunt.normalize_listing({
            **base, "jobUrl": "https://jobs.ashbyhq.com/ziina/123",
        }, source)
        self.assertEqual(
            job_hunt.make_job_key(first["company"], first["title"], first["location"], first["canonical_job_url"]),
            job_hunt.make_job_key(second["company"], second["title"], second["location"], second["canonical_job_url"]),
        )

    def test_greenhouse_and_lever_normalization_regression(self):
        greenhouse = job_hunt.normalize_listing({
            "title": "Account Manager",
            "company_name": "Careem",
            "location": {"name": "Dubai, United Arab Emirates"},
            "absolute_url": "https://boards.greenhouse.io/careem/jobs/123",
            "content": "<p>Client accounts</p>",
        }, {"provider": "greenhouse", "board": "careem", "name": "Careem"})
        lever = job_hunt.normalize_listing({
            "text": "Implementation Specialist",
            "company": "Deliverect",
            "categories": {"location": "Dubai, UAE"},
            "hostedUrl": "https://jobs.lever.co/deliverect/123",
            "applyUrl": "https://jobs.lever.co/deliverect/123/apply",
            "descriptionPlain": "Customer implementation",
        }, {"provider": "lever", "board": "deliverect", "name": "Deliverect"})
        self.assertEqual(greenhouse["location"], "Dubai, United Arab Emirates")
        self.assertEqual(greenhouse["canonical_job_url"], "https://boards.greenhouse.io/careem/jobs/123")
        self.assertEqual(lever["application_url"], "https://jobs.lever.co/deliverect/123/apply")
        self.assertEqual(lever["description"], "Customer implementation")

    def test_fallback_deduplication_without_url(self):
        self.assertEqual(job_hunt.make_job_key("Careem", "Account Manager", "Dubai", None), job_hunt.make_job_key(" careem ", "ACCOUNT MANAGER", " DUBAI ", ""))

    def test_salary_monthly_annual_unknown_and_foreign(self):
        monthly = job_hunt.assess_salary("AED 6,500-8,000 per month", "", 6500)
        annual = job_hunt.assess_salary("AED 78,000-96,000 per annum", "", 6500)
        structured_annual = job_hunt.assess_salary({"min": 78000, "max": 96000, "currency": "AED", "interval": "year"}, "", 6500)
        under = job_hunt.assess_salary("AED 72,000 per year", "", 6500)
        ambiguous = job_hunt.assess_salary("AED 80,000", "", 6500)
        foreign = job_hunt.assess_salary("USD 90,000 per year", "", 6500)
        undisclosed = job_hunt.assess_salary(None, "Competitive salary", 6500)
        self.assertEqual(monthly["assessment"], "meets AED 6,500/month target")
        self.assertEqual(annual["monthly_min_aed"], 6500)
        self.assertEqual(structured_annual["monthly_min_aed"], 6500)
        self.assertTrue(under["exclude"])
        self.assertEqual(ambiguous["period"], "ambiguous")
        self.assertFalse(foreign["exclude"])
        self.assertIn("verification", foreign["assessment"])
        self.assertIn("undisclosed", undisclosed["assessment"])

    def test_salary_range_crossing_target_needs_review(self):
        salary = job_hunt.assess_salary("AED 6,000-8,000 per month", "", 6500)
        self.assertFalse(salary["exclude"])
        self.assertIn("crosses", salary["assessment"])

    def test_uae_only_and_city_priorities(self):
        self.assertEqual(job_hunt.classify_location("Dubai, United Arab Emirates"), ("Dubai", 1))
        self.assertEqual(job_hunt.classify_location("Sharjah, UAE"), ("Sharjah", 2))
        self.assertEqual(job_hunt.classify_location("Abu Dhabi, UAE"), ("Abu Dhabi", 3))
        self.assertEqual(job_hunt.classify_location("Ajman, UAE"), ("Ajman", 4))
        self.assertIsNone(job_hunt.classify_location("Bengaluru, India"))
        self.assertIsNone(job_hunt.classify_location("Remote"))

    def test_relevance_uses_title_evidence_and_years(self):
        job = {"title": "Product Operations Associate", "description": "3+ years in product operations. Own UAT, feature definition, workflows, dashboards, and Excel reporting."}
        score, reasons, exclusion = job_hunt.score_job(job, self.config, self.profile)
        self.assertFalse(exclusion)
        self.assertGreaterEqual(score, self.config["review_threshold"])
        self.assertTrue(any("documented evidence overlaps" in reason for reason in reasons))

    def test_unrelated_and_senior_roles_are_excluded(self):
        unrelated, _, reason = job_hunt.score_job({"title": "Software Engineer", "description": "Python backend development"}, self.config, self.profile)
        senior, _, senior_reason = job_hunt.score_job({"title": "Senior Product Operations Specialist", "description": "Product operations"}, self.config, self.profile)
        self.assertEqual(unrelated, 0)
        self.assertIn("outside", reason)
        self.assertEqual(senior, 0)
        self.assertIn("seniority", senior_reason)

    def test_expired_listing_exclusion(self):
        self.assertTrue(job_hunt.listing_is_expired({"application_deadline": "2025-01-01"}, date(2026, 10, 13)))
        self.assertFalse(job_hunt.listing_is_expired({"application_deadline": "2027-01-01"}, date(2026, 10, 13)))

    def test_upsert_preserves_status_and_date_found(self):
        old = {"job_key": "same", "company": "Careem", "role": "Account Manager", "location": "Dubai", "date_found": "2026-01-01", "status": "Applied", "next_action": "Follow up next week"}
        new = {**old, "date_found": "2026-10-13", "status": "New", "next_action": "Review", "match_score": 72}
        rows = job_hunt.upsert_tracker([old], [new])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["date_found"], "2026-01-01")
        self.assertEqual(rows[0]["status"], "Applied")
        self.assertEqual(rows[0]["next_action"], "Follow up next week")
        self.assertEqual(rows[0]["match_score"], 72)

    def test_city_priority_is_configurable_and_draft_keeps_experience_context(self):
        self.assertEqual(job_hunt.classify_location("Dubai, UAE", ["Abu Dhabi", "Dubai", "Other UAE"]), ("Dubai", 2))
        resume, outreach = job_hunt.build_drafts(
            {"role": "Product Operations Associate", "company": "Example Co"}, self.profile, self.config
        )
        self.assertIn("Accruon Consultant L.L.C.", resume)
        self.assertIn("Dec 2023 - Jan 2026", resume)
        self.assertIn("Cleantel Services", resume)
        self.assertIn("Draft for review only - not sent.", outreach)

    def test_review_draft_trigger_is_thresholded_and_not_repeated(self):
        self.assertTrue(job_hunt.should_generate_review_draft({"match_score": 55}, None, 55))
        self.assertFalse(job_hunt.should_generate_review_draft({"match_score": 54}, None, 55))
        previous = {"resume_draft": "resume.md", "outreach_draft": "outreach.md"}
        self.assertFalse(job_hunt.should_generate_review_draft({"match_score": 90}, previous, 55))

    def test_account_draft_prioritizes_client_evidence(self):
        resume, outreach = job_hunt.build_drafts(
            {"role": "Account Manager I", "company": "Example Co"}, self.profile, self.config
        )
        self.assertIn("SaaS onboarding", resume)
        self.assertIn("followed up with clients on gaps", outreach)
        self.assertIn("one contract renewal", outreach)
        self.assertNotIn("Built client master data", outreach)


if __name__ == "__main__":
    unittest.main()